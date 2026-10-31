"""Read supplier constants only; user torque/speed assumptions remain authoritative."""
import hashlib
import json
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = Path(r'C:\Users\rshan\.codex\codex-remote-attachments\01a08b12-e85a-7f73-8677-315c708cd9f7\78C25BDB-3395-4655-8BFF-24E072FD48AA')
out = {'current_convention': 'Supplier sheets do not identify phase RMS, phase peak, Iq, or DC-bus current. T/Kt is an approximate supplier-convention current, not pack current.', 'sources': [], 'variants': {}}
for filename, columns, rows in [
    ('2-Exoskeleton-Motor-Selection-Table-2026.xlsx', {'WGG6010-8': 3, 'WGG7010-8': 6}, {'voltage_range':2,'rated_voltage_V':3,'winding_resistance_ohm':11,'output_Kv_rpm_V':13,'output_Kt_Nm_A':14,'internal_reduction':16}),
    ('1-Planetary-Reduction-Motor-Selection-Table.xlsx', {'GIM6010-8_24V':13,'GIM6010-8_48V':14}, {'voltage_range':5,'rated_voltage_V':3,'phase_to_phase_resistance_ohm':13,'output_Kv_rpm_V':15,'output_Kt_Nm_A':16,'internal_reduction':18}),
]:
    path = ROOT/filename
    sheet = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    out['sources'].append({'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':sheet.title})
    for name,col in columns.items():
        out['variants'][name] = {k:sheet.cell(row,col).value for k,row in rows.items()}
        out['variants'][name]['source'] = filename
        out['variants'][name]['sheet'] = sheet.title
        out['variants'][name]['cells'] = {k:sheet.cell(row,col).coordinate for k,row in rows.items()}
out['ignored_rows'] = 'Supplier maximum/rated torque and maximum/rated speed are not used to replace user assumptions.'
(HERE/'assets/motor_constants_20260930.json').write_text(json.dumps(out,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(out['variants'],indent=2,ensure_ascii=False))
